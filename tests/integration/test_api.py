"""API contract used by the reviewer UI. Runs against a temporary workspace; no network."""

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from kaizen.api.app import create_app
from kaizen.datasets.build import build_golden
from kaizen.workspace import Workspace

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DHARMA = "dharma.reddy@bd.com"
HEMANT = "hemant@bd.com"
PASSWORD = "review-the-boms-2026"


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    root = tmp_path_factory.mktemp("api")
    golden = build_golden(root / "golden")
    ws = Workspace(root / "ws")
    client = TestClient(create_app(ws))
    # Reviewer identity is a verified BD address now, so the shared client signs in as the facilitator.
    sign_in(client, DHARMA, 1)
    r = client.post("/api/runs/from-path", json={"path": str(golden)})
    assert r.status_code == 200, r.text
    return client, golden, ws, r.json()["run_id"]


def test_health_and_run_listing(env):
    client, golden, ws, run_id = env
    assert client.get("/api/health").json()["status"] == "ok"
    runs = client.get("/api/runs").json()
    assert any(x["run_id"] == run_id for x in runs)
    summary = client.get(f"/api/runs/{run_id}").json()
    assert summary["skus"] == 10 and summary["documents"] == 38 and summary["rows"] > 1000
    assert summary["counts"]["EXACT"] > 0 and summary["needs_validation"] > 0 and summary["auto_cleared"] > 0
    assert any(c["status"] == "MISSING_BOM" for c in summary["coverage"])
    assert "ENGINE_RECOMMENDED" in summary["state_counts"]
    assert summary["capabilities"] and summary["terminology_version"]


def test_results_queue_ordering_and_filters(env):
    client, golden, ws, run_id = env
    queue = client.get(f"/api/runs/{run_id}/results", params={"needs_validation": "true"}).json()
    assert queue["total"] > 0
    assert queue["rows"][0]["engine"]["severity"] == "BLOCKER"
    sevs = [r["engine"]["severity"] or "" for r in queue["rows"]]
    order = {"BLOCKER": 0, "MAJOR": 1, "MINOR": 2, "INFO": 3, "": 4}
    assert all(order[a] <= order[b] for a, b in zip(sevs, sevs[1:]))
    only = client.get(f"/api/runs/{run_id}/results", params={"check": "PCO_BOM", "sku": "1395108QNS", "discrepancy": "PCO_CHANGE_NOT_APPLIED"}).json()
    assert only["total"] == 1 and only["rows"][0]["check"] == "PCO_BOM"
    cls = client.get(f"/api/runs/{run_id}/results", params={"classification": "MISMATCH", "check": "BOM_LABEL"}).json()
    assert cls["total"] >= 3 and all(r["engine"]["classification"] == "MISMATCH" for r in cls["rows"])


def test_row_detail_with_evidence_and_page_image(env):
    client, golden, ws, run_id = env
    row = client.get(f"/api/runs/{run_id}/results", params={"check": "BOM_LABEL", "sku": "1295108FNS", "classification": "MISMATCH"}).json()["rows"][0]
    detail = client.get(f"/api/runs/{run_id}/results/{row['row_id']}").json()
    assert detail["result"]["row_id"] == row["row_id"]
    a, b = detail["evidence"]["a"], detail["evidence"]["b"]
    assert a["doc_id"] and a["page"] == 2 and a["bbox"] and a["locator"] and a["raw_text"]
    assert b["doc_id"] and b["page"] == 1 and b["bbox"]
    assert detail["history"][0]["event"] == "engine"
    img = client.get(f"/api/runs/{run_id}/documents/{a['doc_id']}/pages/{a['page']}", params={"highlight": row["row_id"]})
    assert img.status_code == 200 and img.headers["content-type"] == "image/png" and len(img.content) > 5000
    items = client.get(f"/api/runs/{run_id}/documents/{b['doc_id']}/items").json()
    assert items["doc_type"] == "LABEL" and len(items["items"]) >= 30 and items["items"][0]["bbox"] and items["items"][0]["page"] == 1
    docs = client.get(f"/api/runs/{run_id}/documents").json()
    assert any(d["doc_type"] == "PCO" for d in docs)


def test_decisions_blind_mode_disagreement_and_finalize(env):
    client, golden, ws, run_id = env
    row = client.get(f"/api/runs/{run_id}/results", params={"check": "BOM_LABEL", "sku": "1295108FNS", "classification": "MISMATCH"}).json()["rows"][0]
    rid = row["row_id"]
    r = client.post(f"/api/runs/{run_id}/decisions", json={"row_id": rid, "decision": "CONFIRM_DISCREPANCY", "comment": "qty"})
    assert r.status_code == 200 and r.json()["state"] == "REVIEWER_1_COMPLETE"

    reviewer2 = TestClient(create_app(ws))
    assert sign_in(reviewer2, HEMANT, 2)["blind"] is True
    blind = reviewer2.get(f"/api/runs/{run_id}/results/{rid}").json()
    assert blind["decisions"]["1"] is None and blind["result"]["classification"] == "MISMATCH"
    r = reviewer2.post(f"/api/runs/{run_id}/decisions", json={"row_id": rid, "decision": "ACCEPT"})
    assert r.json()["state"] == "DISAGREEMENT"
    after = reviewer2.get(f"/api/runs/{run_id}/results/{rid}").json()
    assert after["decisions"]["1"]["decision"] == "CONFIRM_DISCREPANCY"

    r = client.post(f"/api/runs/{run_id}/finalize", json={"row_id": rid, "final_decision": "CONFIRM_DISCREPANCY", "note": "meeting"})
    assert r.json()["state"] == "FINALIZED"
    bad = client.post(f"/api/runs/{run_id}/decisions", json={"row_id": rid, "decision": "MAYBE"})
    assert bad.status_code == 400
    n = client.post(f"/api/runs/{run_id}/bulk-accept", json={}).json()["accepted"]
    assert n > 100


def test_save_as_relationship_and_next_run_uses_it(env):
    client, golden, ws, run_id = env
    row = client.get(f"/api/runs/{run_id}/results", params={"check": "BOM_LABEL", "sku": "1295108NS", "classification": "POTENTIAL"}).json()["rows"]
    chlora = next(r for r in row if "CHLORAPREP" in r["engine"]["explanation"].upper())
    r = client.post(f"/api/runs/{run_id}/relationships/from-row", json={"row_id": chlora["row_id"], "by": "Dharma", "scope": "global", "anchor": True, "notes": "confirmed in review"})
    assert r.status_code == 200, r.text
    rel = r.json()
    assert rel["provenance"] == "learned" and rel["item_anchors"] == ["4440003"]
    assert rel["created_by"] == DHARMA, "the signed-in session names the author, not the payload"
    assert "CHLORAPREP APPLICATOR 3ML" in rel["aliases"]
    r2 = client.post("/api/runs/from-path", json={"path": str(golden / "sku-002")})
    run2 = r2.json()["run_id"]
    rows = client.get(f"/api/runs/{run2}/results", params={"check": "BOM_LABEL", "sku": "1295108FNS"}).json()["rows"]
    again = next(x for x in rows if "CHLORAPREP" in x["engine"]["explanation"].upper())
    assert again["engine"]["classification"] == "EQUIVALENT" and again["engine"]["relationship_id"] == rel["id"]


def test_mining_action_items_verify_close_business_and_exports(env):
    client, golden, ws, run_id = env
    sugg = client.get(f"/api/runs/{run_id}/mining").json()
    assert sugg and sugg[0]["sku_count"] >= 2 and sugg[0]["evidence"]
    s = next(x for x in sugg if "LIDOCAINE" in x["a_text"].upper())
    r = client.post(f"/api/runs/{run_id}/mining/approve", json={"a_key": s["a_key"], "b_key": s["b_key"], "by": "Hemant", "scope": "global"})
    assert r.status_code == 200 and r.json()["id"].startswith("REL-")
    s2 = next(x for x in sugg if "GUIDEWIRE" in x["a_text"].upper())
    assert client.post(f"/api/runs/{run_id}/mining/reject", json={"a_key": s2["a_key"], "b_key": s2["b_key"], "by": "Hemant", "note": "different"}).status_code == 200
    assert not any(x["a_key"] == s2["a_key"] for x in client.get(f"/api/runs/{run_id}/mining").json())

    row = client.get(f"/api/runs/{run_id}/results", params={"check": "PCO_BOM", "sku": "1395108QNS", "discrepancy": "PCO_CHANGE_NOT_APPLIED"}).json()["rows"][0]
    ai = client.post(f"/api/runs/{run_id}/action-items", json={"row_id": row["row_id"], "reviewer": "Dharma", "owner": "R&D"}).json()
    assert ai["id"].startswith("AI-") and ai["status"] == "OPEN"
    assert client.patch(f"/api/action-items/{ai['id']}", json={"status": "IN_PROGRESS", "by": "Hemant"}).json()["status"] == "IN_PROGRESS"
    assert any(x["id"] == ai["id"] for x in client.get("/api/action-items").json())
    r3 = client.post("/api/runs/from-path", json={"path": str(golden / "sku-002")}).json()["run_id"]
    vc = client.post(f"/api/runs/{r3}/verify-and-close").json()
    assert ai["id"] in vc["not_covered"]

    bc = client.get(f"/api/runs/{run_id}/business-case").json()
    assert bc["skus"] == 10 and "annual_savings" in bc and bc["assumptions"]
    bc2 = client.get(f"/api/runs/{run_id}/business-case", params={"minutes_per_validation_row": 5}).json()
    assert bc2["estimated_minutes_per_sku"] > bc["estimated_minutes_per_sku"]

    x = client.get(f"/api/runs/{run_id}/export.xlsx")
    assert x.status_code == 200 and "spreadsheet" in x.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert "Action_Items" in wb.sheetnames
    docs = client.get(f"/api/runs/{run_id}/documents").json()
    bom = next(d for d in docs if d["doc_type"] == "BOM" and d["sku"] == "1295108FNS")
    pdf = client.get(f"/api/runs/{run_id}/annotated-bom/{bom['id']}")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"


def test_terminology_endpoints(env):
    client, golden, ws, run_id = env
    rels = client.get("/api/terminology").json()
    assert any(r["id"] == "REL-001" and "usage" in r for r in rels)
    r = client.post("/api/terminology", json={"canonical": "Vessel Dilator", "aliases": ["DILATOR VESSEL"], "scope": "global", "by": "rahul", "notes": "api"})
    rid = r.json()["id"]
    assert client.put(f"/api/terminology/{rid}", json={"aliases": ["DILATOR VESSEL", "VESSEL DILATOR"], "by": "rahul", "note": "alias"}).json()["version"] == 2
    assert len(client.get(f"/api/terminology/{rid}/history").json()) == 2
    assert client.post(f"/api/terminology/{rid}/deactivate", json={"by": "rahul"}).json()["active"] is False
    assert client.get("/api/terminology", params={"search": "vessel", "all": "true"}).json()
    x = client.get("/api/terminology/export.xlsx")
    assert x.status_code == 200
    up = client.post("/api/terminology/import", files={"file": ("rels.xlsx", x.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data={"by": "rahul"})
    assert up.status_code == 200 and "unchanged" in up.json()["summary"]


def test_upload_and_demo(env):
    client, golden, ws, run_id = env
    files = []
    for p in sorted((golden / "sku-001").iterdir()):
        files.append(("files", (f"sku-001/{p.name}", p.read_bytes(), "application/octet-stream")))
    r = client.post("/api/runs/upload", files=files)
    assert r.status_code == 200, r.text
    assert r.json()["skus"] == 1
    d = client.post("/api/demo/load")
    assert d.status_code == 200 and d.json()["skus"] >= 8


# ---- reviewer sessions ---------------------------------------------------------------------------
# Identity and blind mode are server-side. These tests use their own clients so the module fixture's
# cookie jar is not disturbed.


def _fresh(ws):
    return TestClient(create_app(ws))


def sign_in(client, email: str, slot: int, blind=None) -> dict:
    """Through the front door. The account is created once; 409 afterwards just means it already is."""
    assert client.post("/api/auth/signup", json={"email": email, "password": PASSWORD}).status_code in (200, 409)
    body = {"email": email, "password": PASSWORD, "slot": slot}
    if blind is not None:
        body["blind"] = blind
    r = client.post("/api/auth/signin", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def test_session_lifecycle_and_cookie(env):
    _, _, ws, _ = env
    c = _fresh(ws)
    assert c.get("/api/sessions/current").json()["session"] is None
    s = sign_in(c, DHARMA, 1)
    assert s["reviewer"] == DHARMA and s["slot"] == 1 and s["blind"] is False
    assert "token" not in s, "the token belongs in the cookie, not the response body"
    assert c.get("/api/sessions/current").json()["session"]["reviewer"] == DHARMA
    assert c.delete("/api/sessions/current").json()["ended"] is True
    assert c.get("/api/sessions/current").json()["session"] is None
    assert c.post("/api/auth/signin", json={"email": DHARMA, "password": PASSWORD, "slot": 7}).status_code == 400


def test_review_endpoints_refuse_an_anonymous_caller(env):
    _, _, ws, run_id = env
    c = _fresh(ws)
    for path in (f"/api/runs/{run_id}/results", f"/api/runs/{run_id}/export.xlsx", "/api/audit"):
        assert c.get(path).status_code == 401, path
    assert c.post(f"/api/runs/{run_id}/decisions", json={"row_id": "x", "decision": "ACCEPT"}).status_code == 401
    assert c.get(f"/api/runs/{run_id}").status_code == 401


def test_blind_mode_is_taken_from_the_session_not_the_query_string(env):
    _, _, ws, run_id = env
    c1 = _fresh(ws)
    sign_in(c1, DHARMA, 1)
    row = c1.get(f"/api/runs/{run_id}/results", params={"check": "BOM_LABEL", "classification": "MISMATCH"}).json()["rows"][0]
    rid = row["row_id"]
    c1.post(f"/api/runs/{run_id}/decisions", json={"row_id": rid, "decision": "OVERRIDE", "override_classification": "EQUIVALENT", "comment": "same part"})

    c2 = _fresh(ws)
    s = sign_in(c2, HEMANT, 2)
    assert s["blind"] is True and s["blind_review_policy"] == "required"
    # The old client-controlled parameters are ignored.
    page = c2.get(f"/api/runs/{run_id}/results", params={"viewer": 1, "blind": "false", "check": "BOM_LABEL"}).json()
    assert page["viewer"]["slot"] == 2 and page["viewer"]["blind"] is True
    mine = next(r for r in page["rows"] if r["row_id"] == rid)
    assert mine["decisions"]["1"] is None
    assert mine["effective_classification"] == "MISMATCH", "reviewer 1's override leaked to a blind reviewer"
    assert mine["state"] == "ENGINE_RECOMMENDED"

    detail = c2.get(f"/api/runs/{run_id}/results/{rid}", params={"viewer": 1, "blind": "false"}).json()
    assert detail["decisions"]["1"] is None and len(detail["history"]) == 1

    c2.post(f"/api/runs/{run_id}/decisions", json={"row_id": rid, "decision": "ACCEPT"})
    after = c2.get(f"/api/runs/{run_id}/results/{rid}").json()
    assert after["decisions"]["1"]["decision"] == "OVERRIDE" and after["state"] == "DISAGREEMENT"


def test_decision_identity_comes_from_the_session(env):
    _, _, ws, run_id = env
    c = _fresh(ws)
    sign_in(c, HEMANT, 2)
    row = c.get(f"/api/runs/{run_id}/results", params={"check": "BOM_DRAWING"}).json()["rows"][0]
    # A caller claiming a different slot or name is refused / ignored, not believed.
    bad = c.post(f"/api/runs/{run_id}/decisions", json={"row_id": row["row_id"], "slot": 1, "decision": "ACCEPT"})
    assert bad.status_code == 400 and "slot" in bad.json()["detail"].lower()
    ok = c.post(f"/api/runs/{run_id}/decisions", json={"row_id": row["row_id"], "reviewer": "Dharma", "decision": "ACCEPT"}).json()
    assert ok["decisions"]["2"]["reviewer"] == HEMANT and ok["decisions"]["2"]["blind"] is True
    assert "1" not in {k for k, v in ok["decisions"].items() if v}


def test_blind_reviewer_cannot_read_the_audit_log_or_export_the_workbook(env):
    _, _, ws, run_id = env
    c = _fresh(ws)
    sign_in(c, HEMANT, 2)
    for path in ("/api/audit", f"/api/runs/{run_id}/export.xlsx"):
        r = c.get(path)
        assert r.status_code == 403, path
        assert "blind" in r.json()["detail"].lower()
    c2 = _fresh(ws)
    sign_in(c2, DHARMA, 1)
    assert c2.get("/api/audit").status_code == 200
    assert c2.get(f"/api/runs/{run_id}/export.xlsx").status_code == 200


def test_policy_can_be_relaxed_and_then_reviewer_two_may_be_unblinded(env):
    _, _, ws, run_id = env
    from kaizen.review.sessions import POLICY_OPTIONAL, POLICY_REQUIRED, SessionStore

    sessions = SessionStore(ws.db)
    sessions.set_policy(POLICY_OPTIONAL, by="Rahul")
    try:
        c = _fresh(ws)
        assert sign_in(c, HEMANT, 2, blind=False)["blind"] is False
        anon = _fresh(ws)
        assert anon.get(f"/api/runs/{run_id}/results").status_code == 200, "optional policy keeps the anonymous path open"
    finally:
        sessions.set_policy(POLICY_REQUIRED, by="Rahul")


# ---- certificate, diff, optional Excel round-trip -----------------------------------------------


def test_certificate_pdf_for_one_sku_and_for_the_run(env):
    client, _, ws, run_id = env
    one = client.get(f"/api/runs/{run_id}/certificate.pdf", params={"sku": "1295108NS"})
    assert one.status_code == 200 and one.content.startswith(b"%PDF") and "1295108NS" in one.headers["content-disposition"]
    whole = client.get(f"/api/runs/{run_id}/certificate.pdf")
    assert whole.status_code == 200 and len(whole.content) > len(one.content)
    assert client.get(f"/api/runs/{run_id}/certificate.pdf", params={"sku": "nope"}).status_code == 404
    blind = _fresh(ws)
    sign_in(blind, HEMANT, 2)
    assert blind.get(f"/api/runs/{run_id}/certificate.pdf").status_code == 403


def test_diff_against_another_run(env):
    client, _, _, run_id = env
    d = client.get(f"/api/runs/{run_id}/diff", params={"against": run_id}).json()
    assert d["before_run_id"] == run_id and d["counts"]["unchanged"] > 0 and d["counts"]["resolved"] == 0 and d["counts"]["new"] == 0
    assert d["counts"]["still_open"] > 0 and all(r["status"] == "still_open" for r in d["rows"]), "a run compared with itself: open discrepancies stay open"
    assert client.get(f"/api/runs/{run_id}/diff", params={"against": "run-nope"}).status_code == 404


def test_excel_round_trip_through_the_api(env):
    client, _, ws, run_id = env
    row = client.get(f"/api/runs/{run_id}/results", params={"check": "BOM_DRAWING", "classification": "MISSING", "limit": 1}).json()["rows"][0]
    xlsx = client.get(f"/api/runs/{run_id}/export.xlsx").content
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    sh = wb["BOM_Drawing"]
    cols = {c.value: i + 1 for i, c in enumerate(sh[1])}
    for r in range(2, sh.max_row + 1):
        if sh.cell(row=r, column=cols["Row ID"]).value == row["row_id"]:
            sh.cell(row=r, column=cols["Reviewer Decision"], value="CONFIRM_DISCREPANCY")
            sh.cell(row=r, column=cols["Reviewer Comment"], value="decided in Excel")
    buf = io.BytesIO()
    wb.save(buf)
    preview = client.post(f"/api/runs/{run_id}/decisions/import", files={"file": ("report.xlsx", buf.getvalue(), XLSX_MIME)}, data={"dry_run": "true"}).json()
    assert preview["dry_run"] is True and preview["applied"] == [row["row_id"]] and "Dry run" in preview["summary"]
    assert client.get(f"/api/runs/{run_id}/results/{row['row_id']}").json()["decisions"]["1"] is None
    applied = client.post(f"/api/runs/{run_id}/decisions/import", files={"file": ("report.xlsx", buf.getvalue(), XLSX_MIME)}).json()
    assert applied["applied"] == [row["row_id"]]
    d = client.get(f"/api/runs/{run_id}/results/{row['row_id']}").json()["decisions"]["1"]
    assert d["decision"] == "CONFIRM_DISCREPANCY" and d["reviewer"] == DHARMA and d["comment"] == "decided in Excel"
    wrong = client.post(f"/api/runs/{run_id}/decisions/import", files={"file": ("x.xlsx", b"not a workbook", XLSX_MIME)})
    assert wrong.status_code == 400


def test_index_page_is_served_with_no_cache(tmp_path):
    """A rebuilt bundle must show up on the next load: index.html is never cached, hashed assets may be."""
    ui = tmp_path / "dist"
    ui.mkdir()
    (ui / "index.html").write_text("<!doctype html><title>t</title>")
    (ui / "app.js").write_text("// bundle")
    c = TestClient(create_app(Workspace(tmp_path / "ws"), ui_dir=ui))
    assert c.get("/").headers.get("cache-control") == "no-cache"
    assert c.get("/index.html").headers.get("cache-control") == "no-cache"
    assert c.get("/app.js").status_code == 200 and c.get("/app.js").headers.get("cache-control") is None
